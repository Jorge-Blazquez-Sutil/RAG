"""Extractor de DOCX basado en python-docx (licencia MIT)."""

from pathlib import Path

from docx import Document
from docx.opc.exceptions import OpcError, PackageNotFoundError

from app.ingestion.errors import CorruptDocumentError, EmptyTextLayerError
from app.ingestion.extractors.base import DocumentExtractor, normalize_text
from app.schemas.document import DocumentMetadata, ExtractedDocument, PageContent


class DocxExtractor(DocumentExtractor):
    """Extrae los párrafos del cuerpo del documento.

    Nota sobre paginación: un DOCX no tiene páginas fijas — la paginación la calcula
    el procesador de texto al renderizar. Por eso el documento se devuelve como una
    única página; la cita apuntará al documento, no a una página concreta.
    """

    extensions = frozenset({".docx"})

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            document = Document(str(path))
        except (PackageNotFoundError, OpcError, ValueError) as exc:
            raise CorruptDocumentError(path.name, str(exc)) from exc

        text = normalize_text("\n".join(p.text for p in document.paragraphs))
        pages = [PageContent(page_number=1, text=text)]

        extracted = ExtractedDocument(
            metadata=self._metadata(document, path, text),
            pages=pages,
        )
        if extracted.is_empty:
            raise EmptyTextLayerError(path.name)
        return extracted

    def _metadata(self, document: object, path: Path, text: str) -> DocumentMetadata:
        props = document.core_properties  # type: ignore[attr-defined]
        return DocumentMetadata(
            filename=path.name,
            extension=path.suffix.lower(),
            title=(props.title or None),
            author=(props.author or None),
            created_at=props.created,
            page_count=1,
            char_count=len(text),
        )
